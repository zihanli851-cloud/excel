from __future__ import annotations

import hashlib
import logging
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.models.project import Project
from app.schemas.import_result import ImportResult, ImportWarning, SheetImportStat
from app.services.amount_parser import normalize_text, parse_amount
from app.services.excel_utils import (
    find_header_row,
    is_empty_values,
    is_yellow_row,
    parse_date,
    parse_sheet_year,
)
from app.services.invalid_classifier import classify_invalid_reason
from app.services.project_repository import ProjectRepository

logger = logging.getLogger(__name__)


class ExcelImporter:
    def __init__(self, db: Session):
        self.repository = ProjectRepository(db)
        self.settings = get_settings()

    def import_file(self, file_path: Path) -> ImportResult:
        workbook = load_workbook(file_path, data_only=False)
        result = ImportResult(file_name=file_path.name)
        all_projects: list[Project] = []

        for worksheet in workbook.worksheets:
            sheet_year = parse_sheet_year(worksheet.title)
            if sheet_year is None:
                result.warnings.append(
                    ImportWarning(sheet=worksheet.title, message="Sheet 名无法解析年份，已跳过")
                )
                continue

            try:
                header_row, header_map, extra_headers = find_header_row(worksheet)
            except ValueError as exc:
                result.warnings.append(ImportWarning(sheet=worksheet.title, message=str(exc)))
                continue

            if extra_headers:
                result.warnings.append(
                    ImportWarning(
                        sheet=worksheet.title,
                        row=header_row,
                        message=f"发现额外表头，已忽略: {', '.join(extra_headers)}",
                    )
                )

            sheet_stat = SheetImportStat(sheet_name=worksheet.title, sheet_year=sheet_year)
            result.sheet_stats.append(sheet_stat)

            data_start = header_row + 1
            for row_index in range(data_start, worksheet.max_row + 1):
                row_cells = list(worksheet[row_index])
                values = {field: worksheet.cell(row=row_index, column=column).value for field, column in header_map.items()}
                core_values = list(values.values())

                if is_empty_values(core_values):
                    result.skipped_rows += 1
                    continue

                sheet_stat.total_rows += 1
                result.total_rows += 1

                is_yellow = is_yellow_row(row_cells)
                row_text = " ".join(normalize_text(value) for value in core_values if value is not None)
                invalid_reason = classify_invalid_reason(
                    f"{values.get('bid_amount') or ''} {row_text}",
                    is_yellow,
                )
                is_invalid = invalid_reason is not None

                commission_num = parse_amount(values.get("commission_amount"))
                max_price_num = parse_amount(values.get("max_price"))
                bid_amount_num = parse_amount(values.get("bid_amount"))
                self._append_amount_warning(result, worksheet.title, row_index, "commission_amount", values.get("commission_amount"), commission_num)
                self._append_amount_warning(result, worksheet.title, row_index, "max_price", values.get("max_price"), max_price_num)
                self._append_amount_warning(result, worksheet.title, row_index, "bid_amount", values.get("bid_amount"), bid_amount_num)

                project = Project(
                    seq_no=_parse_int(values.get("seq_no")),
                    project_name=normalize_text(values.get("project_name")) or "(未命名项目)",
                    project_code=normalize_text(values.get("project_code")) or None,
                    purchaser=normalize_text(values.get("purchaser")) or None,
                    bid_open_date=parse_date(values.get("bid_open_date")),
                    commission_amount=_text_or_none(values.get("commission_amount")),
                    commission_num=commission_num,
                    max_price=_text_or_none(values.get("max_price")),
                    max_price_num=max_price_num,
                    bid_amount=_text_or_none(values.get("bid_amount")),
                    bid_amount_num=bid_amount_num,
                    bid_amount_detail=_text_or_none(values.get("bid_amount_detail")),
                    sheet_year=sheet_year,
                    is_invalid=is_invalid,
                    invalid_reason=invalid_reason,
                    source_file=file_path.name,
                    source_sheet=worksheet.title,
                    source_row=row_index,
                    row_hash=_make_row_hash(sheet_year, worksheet.title, row_index, values),
                )
                all_projects.append(project)
                sheet_stat.imported_rows += 1
                result.imported_rows += 1
                if is_invalid:
                    sheet_stat.invalid_rows += 1
                    result.invalid_rows += 1
                else:
                    result.valid_rows += 1

        written = 0
        for batch in _chunks(all_projects, self.settings.import_batch_size):
            written += self.repository.bulk_upsert(batch)
        if written:
            logger.info("Imported %s rows from %s", written, file_path)

        return result

    @staticmethod
    def _append_amount_warning(
        result: ImportResult,
        sheet: str,
        row: int,
        field: str,
        raw_value: Any,
        parsed_value: object | None,
    ) -> None:
        text = normalize_text(raw_value)
        if text and parsed_value is None and not _is_known_non_numeric_amount(text):
            result.warnings.append(
                ImportWarning(sheet=sheet, row=row, field=field, message=f"金额字段未解析为数字: {text}")
            )


def _text_or_none(value: Any) -> str | None:
    text = normalize_text(value)
    return text or None


def _parse_int(value: Any) -> int | None:
    text = normalize_text(value)
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def _make_row_hash(sheet_year: int, sheet_name: str, row_index: int, values: dict[str, Any]) -> str:
    parts = [
        str(sheet_year),
        sheet_name,
        str(values.get("seq_no") or ""),
        normalize_text(values.get("project_code")),
        normalize_text(values.get("project_name")),
        normalize_text(values.get("purchaser")),
        normalize_text(values.get("bid_open_date")),
        str(row_index),
    ]
    return hashlib.sha256("|".join(parts).encode("utf-8")).hexdigest()


def _chunks(projects: list[Project], batch_size: int) -> list[list[Project]]:
    safe_size = max(batch_size, 1)
    return [projects[index : index + safe_size] for index in range(0, len(projects), safe_size)]


def _is_known_non_numeric_amount(text: str) -> bool:
    if text in {"无", "暂无", "无金额"}:
        return True
    known_terms = ("单价限价", "废标", "终止采购", "采购任务取消", "放弃中标", "无法履约")
    return any(term in text for term in known_terms)
