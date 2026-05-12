from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.config import get_settings
from app.models.project import Project
from app.schemas.export import ExportResult

HEADERS = [
    "序号",
    "项目名称",
    "项目编号",
    "采购人",
    "开标时间",
    "委托金额(万元)合计",
    "最高限价(万元)（分包）",
    "中标金额(万元)合计",
    "中标金额(万元)（分包）",
    "无效原因",
    "所属年度",
]

INVALID_REASON_LABELS = {
    "BID_FAILED": "废标",
    "TERMINATED": "终止采购",
    "CANCELLED": "采购取消",
    "WINNER_QUIT": "中标人放弃",
    "BREACH": "履约失败",
    "OTHER": "其他",
}


def export_projects(projects: Iterable[Project], file_name: str | None = None) -> ExportResult:
    settings = get_settings()
    settings.export_path.mkdir(parents=True, exist_ok=True)

    rows = list(projects)
    if file_name is None:
        file_name = f"查询结果_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    if not file_name.lower().endswith(".xlsx"):
        file_name = f"{file_name}.xlsx"

    target = settings.export_path / Path(file_name).name
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "查询结果"
    worksheet.append(HEADERS)

    header_font = Font(bold=True)
    yellow_fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    for cell in worksheet[1]:
        cell.font = header_font

    for project in rows:
        worksheet.append(_project_to_row(project))
        if project.is_invalid:
            for cell in worksheet[worksheet.max_row]:
                cell.fill = yellow_fill

    worksheet.freeze_panes = "A2"
    _auto_width(worksheet)
    workbook.save(target)
    return ExportResult(file_name=target.name, file_path=target, row_count=len(rows))


def _project_to_row(project: Project) -> list[object | None]:
    return [
        project.seq_no,
        project.project_name,
        project.project_code,
        project.purchaser,
        _format_date(project.bid_open_date),
        project.commission_amount,
        project.max_price,
        project.bid_amount,
        project.bid_amount_detail,
        INVALID_REASON_LABELS.get(project.invalid_reason or "", project.invalid_reason),
        project.sheet_year,
    ]


def _format_date(value: date | None) -> str | None:
    if value is None:
        return None
    return value.isoformat()


def _auto_width(worksheet) -> None:
    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 10), 40)
