from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from app.services.excel_utils import find_header_row, is_yellow_cell, parse_sheet_year


def test_parse_sheet_year() -> None:
    assert parse_sheet_year("2023年") == 2023
    assert parse_sheet_year("项目清单2026") == 2026
    assert parse_sheet_year("说明") is None


def test_find_header_row_with_normalized_headers() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["说明"])
    worksheet.append([
        "序号",
        "项目名称",
        "项目编号",
        "采购人",
        "开标时间",
        "委托金额(万元)合计",
        "最高限价(万元)（分包）",
        "中标金额(万元)合计",
        "中标金额(万元)（分包）",
    ])
    row_index, mapping, extra_headers = find_header_row(worksheet)
    assert row_index == 2
    assert mapping["project_name"] == 2
    assert mapping["max_price"] == 7
    assert extra_headers == []


def test_find_header_row_reports_extra_headers() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append([
        "序号",
        "项目名称",
        "项目编号",
        "采购人",
        "开标时间",
        "委托金额(万元)合计",
        "最高限价(万元)（分包）",
        "中标金额(万元)合计",
        "中标金额(万元)（分包）",
        "备注",
    ])
    _, _, extra_headers = find_header_row(worksheet)
    assert extra_headers == ["备注"]


def test_is_yellow_cell() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    cell = worksheet["A1"]
    cell.fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    assert is_yellow_cell(cell) is True
