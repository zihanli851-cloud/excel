from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker

from app.core.database import Base
from app.models.project import Project
from app.services.excel_importer import ExcelImporter


def test_import_file_reads_projects() -> None:
    artifact_dir = Path("test_artifacts/importer")
    artifact_dir.mkdir(parents=True, exist_ok=True)
    excel_path = artifact_dir / "projects.xlsx"
    _make_workbook(excel_path)

    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False, expire_on_commit=False)

    with SessionLocal() as db:
        result = ExcelImporter(db).import_file(excel_path)
        projects = db.scalars(select(Project).order_by(Project.seq_no)).all()

    assert result.total_rows == 2
    assert result.imported_rows == 2
    assert result.invalid_rows == 1
    assert any("额外表头" in warning.message for warning in result.warnings)
    assert any(warning.field == "commission_amount" for warning in result.warnings)
    assert len(projects) == 2
    assert projects[0].project_name == "正常项目"
    assert str(projects[0].commission_num) == "360.2800"
    assert projects[1].is_invalid is True
    assert projects[1].invalid_reason == "BID_FAILED"


def _make_workbook(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "2026年"
    worksheet.append([
        "序号",
        "项目名称",
        "项目编号",
        "采购人",
        "开标时间",
        "委托金额(万元)合计",
        "最高限价(万元)（分包）",
        "中标金额(万元)合计",
        "中标金额(万元)（分包）",
        "额外字段",
    ])
    worksheet.append([1, "正常项目", "P20260001", "采购人A", "2026-05-01", "360.28万元", "500", "320", "分包1", "忽略"])
    worksheet.append([2, "废标项目", "P20260002", "采购人B", "2026-05-02", "无法确定", "单价限价", "投标人不足三家，废标", "", "忽略"])
    yellow = PatternFill(fill_type="solid", fgColor="FFFF00")
    for cell in worksheet[3]:
        cell.fill = yellow
    workbook.save(path)
